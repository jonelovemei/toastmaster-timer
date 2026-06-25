"""Toastmaster Timer - web edition (PyScript / Pyodide).

Same logic as the desktop tool, running entirely in the browser:
- timer with green/yellow/red/max-red phases + buzzer (Web Audio)
- bilingual UI (chosen on the login page)
- BLE light signalling via Web Bluetooth (HM-10 / HM-19: service 0xFFE0, char 0xFFE1)
- template-based report export (XLSX via openpyxl, plus CSV / TXT / PDF)
"""
import asyncio
import io
from copy import copy
from datetime import datetime

from pyscript import document, window
from pyodide.ffi import create_proxy
from openpyxl import load_workbook

TEMPLATE_FILENAME = 'TimeReport_Template.xlsx'

# ----------------------------------------------------------------------------
# Report template layout (mirrors the desktop version).
# ----------------------------------------------------------------------------
SECTION_LAYOUT = {
    'table_topic': {
        'data_start': 12, 'data_end': 16, 'has_proj': False,
        'left':  {'no': 'A', 'name': 'B', 'proj': None, 'time': 'G', 'qual': 'I'},
        'right': {'no': 'L', 'name': 'M', 'proj': None, 'time': 'Q', 'qual': 'S'},
    },
    'prepared': {
        'data_start': 20, 'data_end': 24, 'has_proj': True,
        'left':  {'no': 'A', 'name': 'B', 'proj': 'E', 'time': 'G', 'qual': 'I'},
        'right': {'no': 'L', 'name': 'M', 'proj': 'O', 'time': 'Q', 'qual': 'S'},
    },
    'evaluator': {
        'data_start': 28, 'data_end': 32, 'has_proj': False,
        'left':  {'no': 'A', 'name': 'B', 'proj': None, 'time': 'G', 'qual': 'I'},
        'right': {'no': 'L', 'name': 'M', 'proj': None, 'time': 'Q', 'qual': 'S'},
    },
}

SECTION_TIMING = {
    'table_topic': ("TABLE TOPIC (Timing: 1 to 2 min)", "即兴演讲 桌话题 (时长：1 至 2 分钟)"),
    'prepared': ("PREPARED SPEECH (Timing: P1: 4 to 6 min, P2 to P9: 5 to 7 min, P10: 8 to 10 min)",
                 "备稿演讲 (时长：P1：4 至 6 分钟，P2 至 P9：5 至 7 分钟，P10：8 至 10 分钟)"),
    'evaluator': ("SPEECH EVALUATOR (Timing: 2 to 3 min)", "演讲点评 (时长：2 至 3 分钟)"),
}

SECTION_ORDER = ['table_topic', 'prepared', 'evaluator']


def _copy_cell_style(ws, src_row, src_col, dst_row, dst_col):
    s = ws.cell(row=src_row, column=src_col)
    d = ws.cell(row=dst_row, column=dst_col)
    if s.has_style:
        d.font = copy(s.font)
        d.border = copy(s.border)
        d.fill = copy(s.fill)
        d.alignment = copy(s.alignment)
        d.number_format = s.number_format
        d.protection = copy(s.protection)


def _row_merges(ws, row):
    return [(m.min_col, m.max_col) for m in ws.merged_cells.ranges
            if m.min_row == row and m.max_row == row]


def _shift_rows_down(ws, first_row, count, max_col=28):
    last = ws.max_row
    moved = [(m.min_row, m.min_col, m.max_row, m.max_col)
             for m in list(ws.merged_cells.ranges) if m.min_row >= first_row]
    for (mr, mc, xr, xc) in moved:
        ws.unmerge_cells(start_row=mr, start_column=mc, end_row=xr, end_column=xc)
    for r in range(last, first_row - 1, -1):
        for c in range(1, max_col + 1):
            s = ws.cell(row=r, column=c)
            d = ws.cell(row=r + count, column=c)
            d.value = s.value
            _copy_cell_style(ws, r, c, r + count, c)
        if r in ws.row_dimensions:
            ws.row_dimensions[r + count].height = ws.row_dimensions[r].height
    for r in range(first_row, first_row + count):
        for c in range(1, max_col + 1):
            ws.cell(row=r, column=c).value = None
    for (mr, mc, xr, xc) in moved:
        ws.merge_cells(start_row=mr + count, start_column=mc,
                       end_row=xr + count, end_column=xc)


def _clone_row(ws, src_row, dst_row, max_col=28):
    for c in range(1, max_col + 1):
        ws.cell(row=dst_row, column=c).value = None
        _copy_cell_style(ws, src_row, c, dst_row, c)
    for (mn, mx) in _row_merges(ws, src_row):
        ws.merge_cells(start_row=dst_row, start_column=mn, end_row=dst_row, end_column=mx)
    if src_row in ws.row_dimensions:
        ws.row_dimensions[dst_row].height = ws.row_dimensions[src_row].height


def _put_record(ws, row, side, no, rec, has_proj):
    ws[f"{side['no']}{row}"] = no
    ws[f"{side['name']}{row}"] = rec['name']
    if has_proj and side['proj']:
        ws[f"{side['proj']}{row}"] = rec.get('proj', '')
    ws[f"{side['time']}{row}"] = rec.get('time', '')
    ws[f"{side['qual']}{row}"] = rec.get('qual', '')


def _fill_section(ws, sec, entries):
    cfg = SECTION_LAYOUT[sec]
    data_start, data_end = cfg['data_start'], cfg['data_end']
    has_proj = cfg['has_proj']
    base_rows = data_end - data_start + 1
    cap = base_rows * 2
    n = len(entries)
    if n > cap:
        extra_pairs = (n - cap + 1) // 2
        _shift_rows_down(ws, data_end + 1, extra_pairs)
        for k in range(extra_pairs):
            _clone_row(ws, data_end, data_end + 1 + k)
    for idx in range(n):
        rec = entries[idx]
        no = idx + 1
        if idx < base_rows:
            _put_record(ws, data_start + idx, cfg['left'], no, rec, has_proj)
        elif idx < cap:
            _put_record(ws, data_start + (idx - base_rows), cfg['right'], no, rec, has_proj)
        else:
            j = idx - cap
            row = data_end + 1 + (j // 2)
            side = cfg['left'] if j % 2 == 0 else cfg['right']
            _put_record(ws, row, side, no, rec, has_proj)


def populate_template_workbook(wb, sections, header):
    ws = wb['Timer'] if 'Timer' in wb.sheetnames else wb.active
    if header.get('date'):
        ws['I7'] = header['date']
    if header.get('timer_name'):
        ws['R7'] = header['timer_name']
    for sec in reversed(SECTION_ORDER):
        _fill_section(ws, sec, sections.get(sec, []))
    return ws


# ----------------------------------------------------------------------------
# Translations
# ----------------------------------------------------------------------------
TR = {
    'title': ("Toastmaster Timer", "头马演讲计时器"),
    'thresholds_frame': ("Set Time Thresholds (minutes)", "设置时间阈值（分钟）"),
    'green': ("Green (minutes):", "绿灯（分钟）："),
    'yellow': ("Yellow (minutes):", "黄灯（分钟）："),
    'red': ("Red (minutes):", "红灯（分钟）："),
    'max_red': ("Max Red (minutes):", "最大红灯（分钟）："),
    'fast_forward': ("Fast Forward", "快进"),
    'log_end': ("Log End Time", "记录结束时间"),
    'next_speaker': ("Next Speaker", "下一位演讲者"),
    'current_speaker': ("Current Speaker", "当前演讲者"),
    'none': ("None", "无"),
    'roster_frame': ("Roster: Import / Add / Export", "名单：导入 / 添加 / 导出"),
    'name': ("Name:", "姓名："),
    'level': ("Level:", "级别："),
    'add': ("Add", "添加"),
    'remove': ("Remove Selected", "删除所选"),
    'clear': ("Clear", "清空"),
    'export_report_label': ("Export Report:", "导出报告："),
    'export_report_btn': ("Export Report", "导出报告"),
    'apply': ("Apply Times", "应用时间"),
    'start': ("Start", "开始"),
    'pause': ("Pause", "暂停"),
    'reset': ("Reset", "重置"),
    'footer': ("Developed by ZXK from Toastmasters CLUB 6237 | © 2026 All Rights Reserved",
               "由头马 6237 俱乐部 ZXK 开发 | © 2026 版权所有"),
    'bt_title': ("Bluetooth", "蓝牙"),
    'bt_connected': ("Connected: {name}", "已连接：{name}"),
    'bt_connect_failed': ("Bluetooth connection failed: {reason}", "蓝牙连接失败：{reason}"),
    'bt_unsupported': ("This browser does not support Web Bluetooth. Use Chrome or Edge.",
                       "此浏览器不支持 Web Bluetooth，请使用 Chrome 或 Edge。"),
    'bt_disconnected': ("Bluetooth disconnected.", "蓝牙已断开连接。"),
    'bt_lost': ("Bluetooth connection lost. Please reconnect.", "蓝牙连接已断开，请重新连接。"),
    'invalid_order': ("Times must be non-negative and increasing: Green <= Yellow <= Red <= Max Red",
                      "时间须为非负且递增：绿 <= 黄 <= 红 <= 最大红"),
    'invalid_numbers': ("Please enter valid numbers for all time thresholds.", "请为所有时间阈值输入有效数字。"),
    'times_updated': ("Timer thresholds have been updated.", "计时阈值已更新。"),
    'clear_roster_confirm': ("Clear all names from the roster?", "确定要清空名单中的所有姓名吗？"),
    'select_speaker': ("Please select a speaker from the roster first.", "请先在名单中选择一位演讲者。"),
    'logged': ("Logged end time {t} for {n}", "已为 {n} 记录结束时间 {t}"),
    'roster_empty': ("Roster is empty.", "名单为空。"),
    'no_more_speakers': ("No more speakers in roster.", "名单中没有更多演讲者了。"),
    'name_required': ("Please enter a name.", "请输入姓名。"),
    'exported': ("Report exported: {f}", "报告已导出：{f}"),
    'export_fmt_label': ("Export Report:", "导出报告："),
    'meta': ("Timer: {timer}   Date: {date}   YTMC: {ytmc}", "计时员：{timer}   日期：{date}   YTMC：{ytmc}"),
}

# id -> translation key for static UI text
LABELS = {
    'thresholds_title': 'thresholds_frame',
    'lbl_green': 'green', 'lbl_yellow': 'yellow', 'lbl_red': 'red', 'lbl_maxred': 'max_red',
    'ff_btn': 'fast_forward', 'log_btn': 'log_end', 'next_btn': 'next_speaker',
    'roster_title': 'roster_frame', 'lbl_name': 'name', 'lbl_level': 'level',
    'add_btn': 'add', 'remove_btn': 'remove', 'clear_btn': 'clear',
    'lbl_export': 'export_report_label', 'export_btn': 'export_report_btn',
    'apply_btn': 'apply', 'start_btn': 'start', 'pause_btn': 'pause', 'reset_btn': 'reset',
    'footer': 'footer',
}


class App:
    def __init__(self):
        self.lang = 'en'
        self.login = {'timer_name': '', 'date': '', 'ytmc_no': ''}
        self.green_time = 60
        self.yellow_time = 120
        self.red_time = 180
        self.max_red_time = 210
        self.time_elapsed = 0
        self.running = False
        self.current_stage = 'none'
        self.buzzer_enabled = True
        self.ble_connected = False
        self.roster = []            # list of {name, level, end_time}
        self.selected_index = None  # highlighted in list
        self.current_index = None   # current speaker
        self.fast_forward_active = False
        self._proxies = []          # keep event proxies alive
        self._ff_task = None

    # ---- translation ----
    def tr(self, key, **kw):
        pair = TR.get(key)
        if pair is None:
            return key
        txt = pair[0] if self.lang == 'en' else pair[1]
        return txt.format(**kw) if kw else txt

    # ---- small DOM helpers ----
    @staticmethod
    def el(id_):
        return document.getElementById(id_)

    def on(self, id_, evt, handler):
        p = create_proxy(handler)
        self._proxies.append(p)
        self.el(id_).addEventListener(evt, p)

    def toast(self, msg):
        t = self.el('toast')
        t.textContent = msg
        t.classList.add('show')
        async def hide():
            await asyncio.sleep(2.2)
            t.classList.remove('show')
        asyncio.ensure_future(hide())

    # ---- startup / login ----
    def start(self):
        self.el('in_date').value = datetime.now().strftime('%Y-%m-%d')
        self.on('btn_enter', 'click', lambda e: self.do_login())
        for fid in ('in_timer_name', 'in_date', 'in_ytmc'):
            self.on(fid, 'keydown', self._login_enter)
        self.el('in_timer_name').focus()

    def _login_enter(self, e):
        if e.key == 'Enter':
            self.do_login()

    def do_login(self):
        self.login = {
            'timer_name': self.el('in_timer_name').value.strip(),
            'date': self.el('in_date').value.strip(),
            'ytmc_no': self.el('in_ytmc').value.strip(),
        }
        self.lang = 'zh' if self.el('in_lang').value == 'zh' else 'en'
        self.el('login').classList.add('hidden')
        self.el('app').classList.remove('hidden')
        self.build_app()

    def build_app(self):
        document.title = self.tr('title')
        self.apply_language()
        self.update_meta()
        self.wire_app_events()
        self.render_roster()
        self.update_current_speaker_label()
        # start the 1-second timer loop
        asyncio.ensure_future(self._tick_loop())

    def apply_language(self):
        for id_, key in LABELS.items():
            node = self.el(id_)
            if node is not None:
                node.textContent = self.tr(key)
        self.el('export_btn').textContent = self.tr('export_report_btn')
        if not self.ble_connected:
            self.el('ble_btn').textContent = self.tr('bt_title')

    def update_meta(self):
        self.el('meta_line').textContent = self.tr(
            'meta', timer=self.login.get('timer_name', '') or '-',
            date=self.login.get('date', '') or '-',
            ytmc=self.login.get('ytmc_no', '') or '-')

    # ---- event wiring ----
    def wire_app_events(self):
        self.on('apply_btn', 'click', lambda e: self.apply_times())
        self.on('start_btn', 'click', lambda e: self.start_timer())
        self.on('pause_btn', 'click', lambda e: self.pause_timer())
        self.on('reset_btn', 'click', lambda e: self.reset_timer())
        self.on('buzzer_btn', 'click', lambda e: self.toggle_buzzer())
        self.on('ble_btn', 'click', lambda e: asyncio.ensure_future(self.toggle_ble()))
        self.on('add_btn', 'click', lambda e: self.add_name())
        self.on('remove_btn', 'click', lambda e: self.remove_selected())
        self.on('clear_btn', 'click', lambda e: self.clear_roster())
        self.on('log_btn', 'click', lambda e: self.log_end_time())
        self.on('next_btn', 'click', lambda e: self.next_speaker())
        self.on('export_btn', 'click', lambda e: asyncio.ensure_future(self.export_report()))
        self.on('name_in', 'keydown', lambda e: self.add_name() if e.key == 'Enter' else None)
        self.on('level_in', 'keydown', lambda e: self.add_name() if e.key == 'Enter' else None)
        # fast forward: press and hold
        for evt in ('mousedown', 'touchstart'):
            self.on('ff_btn', evt, lambda e: self.start_fast_forward())
        for evt in ('mouseup', 'mouseleave', 'touchend'):
            self.on('ff_btn', evt, lambda e: self.stop_fast_forward())
        # BLE lost callback from JS
        window.onBleLost = create_proxy(self._on_ble_lost)

    # ---- timer ----
    async def _tick_loop(self):
        while True:
            await asyncio.sleep(1)
            if self.running:
                self.time_elapsed += 1
                self.render_time()

    def render_time(self):
        m, s = self.time_elapsed // 60, self.time_elapsed % 60
        self.el('time_display').textContent = f"{m:02d}:{s:02d}"
        self.check_color()

    def check_color(self):
        t = self.time_elapsed
        disp = self.el('time_display')
        if t < self.green_time:
            stage = 'none'; disp.style.background = '#ffffff'; disp.style.color = '#0f172a'
        elif t < self.yellow_time:
            stage = 'green'; disp.style.background = '#16a34a'; disp.style.color = '#ffffff'
        elif t < self.red_time:
            stage = 'yellow'; disp.style.background = '#f59e0b'; disp.style.color = '#0f172a'
        elif t < self.max_red_time:
            stage = 'red'; disp.style.background = '#dc2626'; disp.style.color = '#ffffff'
        else:
            stage = 'maxred'; disp.style.background = '#7f1d1d'; disp.style.color = '#ffffff'
        if stage != self.current_stage:
            self.handle_stage_transition(stage)
            self.current_stage = stage

    def handle_stage_transition(self, new_stage):
        beeps = {'yellow': 1, 'red': 2, 'maxred': 3}
        if new_stage in beeps and self.buzzer_enabled:
            try:
                window.beepTimes(beeps[new_stage])
            except Exception:
                pass
        sig = {'green': 1, 'yellow': 2, 'red': 3, 'maxred': 4}.get(new_stage)
        if sig is not None:
            self.send_signal(sig)

    def start_timer(self):
        self.running = True
        self.el('start_btn').disabled = True
        self.el('pause_btn').disabled = False
        self.el('reset_btn').disabled = False
        self.el('apply_btn').disabled = True

    def pause_timer(self):
        self.running = False
        self.el('start_btn').disabled = False
        self.el('pause_btn').disabled = True
        self.el('reset_btn').disabled = False
        self.el('apply_btn').disabled = False

    def reset_timer(self):
        self.running = False
        self.time_elapsed = 0
        disp = self.el('time_display')
        disp.textContent = '00:00'
        disp.style.background = '#ffffff'; disp.style.color = '#0f172a'
        self.el('start_btn').disabled = False
        self.el('pause_btn').disabled = True
        self.el('reset_btn').disabled = True
        self.el('apply_btn').disabled = False
        self.current_stage = 'none'
        self.send_signal(0)

    # ---- fast forward ----
    def start_fast_forward(self):
        if self.fast_forward_active:
            return
        self.fast_forward_active = True
        self._ff_task = asyncio.ensure_future(self._ff_loop())

    def stop_fast_forward(self):
        self.fast_forward_active = False

    async def _ff_loop(self):
        while self.fast_forward_active:
            self.time_elapsed += 20
            self.render_time()
            await asyncio.sleep(0.2)

    # ---- buzzer ----
    def toggle_buzzer(self):
        self.buzzer_enabled = not self.buzzer_enabled
        btn = self.el('buzzer_btn')
        if self.buzzer_enabled:
            btn.classList.add('on'); btn.innerHTML = '&#128266;'
        else:
            btn.classList.remove('on'); btn.innerHTML = '&#128263;'

    # ---- thresholds ----
    def apply_times(self):
        try:
            g = float(self.el('green_in').value)
            y = float(self.el('yellow_in').value)
            r = float(self.el('red_in').value)
            mr = float(self.el('maxred_in').value)
        except (ValueError, TypeError):
            self.toast(self.tr('invalid_numbers'))
            return
        if not (0 <= g <= y <= r <= mr):
            self.toast(self.tr('invalid_order'))
            return
        self.green_time = int(g * 60)
        self.yellow_time = int(y * 60)
        self.red_time = int(r * 60)
        self.max_red_time = int(mr * 60)
        self.check_color()
        self.toast(self.tr('times_updated'))

    # ---- BLE ----
    async def toggle_ble(self):
        if self.ble_connected:
            window.bleDisconnect()
            self.ble_connected = False
            self._set_ble_button(False)
            self.toast(self.tr('bt_disconnected'))
            return
        if not window.bleSupported():
            self.toast(self.tr('bt_unsupported'))
            return
        try:
            name = await window.bleConnect()
            self.ble_connected = True
            self._set_ble_button(True)
            self.toast(self.tr('bt_connected', name=name))
        except Exception as e:
            self.toast(self.tr('bt_connect_failed', reason=str(e)))

    def _set_ble_button(self, connected):
        btn = self.el('ble_btn')
        if connected:
            btn.classList.add('connected')
            btn.textContent = self.tr('bt_title') + ' ✓'
        else:
            btn.classList.remove('connected')
            btn.textContent = self.tr('bt_title')

    def _on_ble_lost(self, *args):
        self.ble_connected = False
        self._set_ble_button(False)
        self.toast(self.tr('bt_lost'))

    def send_signal(self, byte):
        if not self.ble_connected:
            return
        async def _send():
            try:
                await window.bleSend(byte)
            except Exception:
                self._on_ble_lost()
        asyncio.ensure_future(_send())

    # ---- roster ----
    def add_name(self):
        name = self.el('name_in').value.strip()
        level = self.el('level_in').value.strip()
        if not name:
            self.toast(self.tr('name_required'))
            return
        self.roster.append({'name': name, 'level': level, 'end_time': ''})
        self.el('name_in').value = ''
        self.el('level_in').value = ''
        self.render_roster()
        self.el('name_in').focus()

    def remove_selected(self):
        if self.selected_index is None:
            self.toast(self.tr('select_speaker'))
            return
        del self.roster[self.selected_index]
        if self.current_index is not None:
            if self.current_index == self.selected_index:
                self.current_index = None
            elif self.current_index > self.selected_index:
                self.current_index -= 1
        self.selected_index = None
        self.render_roster()
        self.update_current_speaker_label()

    def clear_roster(self):
        if not self.roster:
            return
        if not window.confirm(self.tr('clear_roster_confirm')):
            return
        self.roster = []
        self.selected_index = None
        self.current_index = None
        self.render_roster()
        self.update_current_speaker_label()

    def roster_display(self, e):
        parts = [e['name']]
        if e['level']:
            parts.append(e['level'])
        if e['end_time']:
            parts.append('End ' + e['end_time'])
        return ' | '.join(parts)

    def render_roster(self):
        ul = self.el('roster_list')
        ul.innerHTML = ''
        for i, e in enumerate(self.roster):
            li = document.createElement('li')
            li.textContent = self.roster_display(e)
            if i == self.selected_index:
                li.classList.add('selected')
            if i == self.current_index:
                li.classList.add('current')
            handler = create_proxy((lambda idx: (lambda ev: self.select_row(idx)))(i))
            self._proxies.append(handler)
            li.addEventListener('click', handler)
            ul.appendChild(li)

    def select_row(self, idx):
        self.selected_index = idx
        self.render_roster()

    def log_end_time(self):
        if self.selected_index is None:
            self.toast(self.tr('select_speaker'))
            return
        t = f"{self.time_elapsed//60:02d}:{self.time_elapsed%60:02d}"
        e = self.roster[self.selected_index]
        e['end_time'] = t
        self.render_roster()
        self.toast(self.tr('logged', t=t, n=e['name']))

    def next_speaker(self):
        if not self.roster:
            self.toast(self.tr('roster_empty'))
            return
        if self.selected_index is not None:
            self.current_index = self.selected_index
        if self.current_index is None:
            nxt = 0
        else:
            # log current speaker end time before moving on
            t = f"{self.time_elapsed//60:02d}:{self.time_elapsed%60:02d}"
            self.roster[self.current_index]['end_time'] = t
            nxt = self.current_index + 1
        if nxt >= len(self.roster):
            self.toast(self.tr('no_more_speakers'))
            self.current_index = None
            self.selected_index = None
            self.render_roster()
            self.update_current_speaker_label()
            return
        self.current_index = nxt
        self.selected_index = nxt
        self.render_roster()
        self.update_current_speaker_label()
        self.reset_timer()
        self.start_timer()

    def update_current_speaker_label(self):
        if self.current_index is not None and 0 <= self.current_index < len(self.roster):
            name = self.roster[self.current_index]['name']
        else:
            name = self.tr('none')
        self.el('current_speaker').textContent = self.tr('current_speaker') + ': ' + name

    # ---- report data ----
    @staticmethod
    def _time_to_seconds(text):
        if not text:
            return None
        text = str(text).strip()
        if ':' not in text:
            return None
        try:
            mm, ss = text.split(':', 1)
            return int(mm) * 60 + int(ss)
        except (ValueError, TypeError):
            return None

    def qualified_status(self, end_time):
        secs = self._time_to_seconds(end_time)
        if secs is None:
            return ''
        return 'Yes' if self.green_time <= secs <= self.max_red_time else 'No'

    def build_report_sections(self):
        sections = {'table_topic': [], 'prepared': [], 'evaluator': []}
        for e in self.roster:
            level = (e.get('level') or '').strip()
            ll = level.lower()
            if 'evaluator' in ll:
                sec = 'evaluator'
            elif 'table topic' in ll:
                sec = 'table_topic'
            else:
                sec = 'prepared'
            end_time = e.get('end_time', '') or ''
            sections[sec].append({
                'name': e.get('name', ''),
                'level': level,
                'time': end_time,
                'qual': self.qualified_status(end_time),
                'proj': level if sec == 'prepared' else '',
            })
        return sections

    def report_meta_lines(self):
        return [
            ('Timer', self.login.get('timer_name', '')),
            ('Date', self.login.get('date', '')),
            ('YTMC No.', self.login.get('ytmc_no', '')),
            ('Generated', datetime.now().strftime('%Y-%m-%d %H:%M:%S')),
            ('Total Speakers', str(len(self.roster))),
        ]

    def section_title(self, sec):
        pair = SECTION_TIMING[sec]
        return pair[0] if self.lang == 'en' else pair[1]

    def section_table(self, sec, records):
        if SECTION_LAYOUT[sec]['has_proj']:
            headers = ['No.', 'Speaker Name', 'Project Level', 'Time Taken', 'Qualified']
            rows = [[str(i + 1), r['name'], r.get('proj', ''), r.get('time', ''), r.get('qual', '')]
                    for i, r in enumerate(records)]
        else:
            headers = ['No.', 'Speaker Name', 'Time Taken', 'Qualified']
            rows = [[str(i + 1), r['name'], r.get('time', ''), r.get('qual', '')]
                    for i, r in enumerate(records)]
        return headers, rows

    # ---- export ----
    async def export_report(self):
        if not self.roster:
            self.toast(self.tr('roster_empty'))
            return
        fmt = self.el('export_fmt').value
        if fmt == 'xlsx':
            self.export_xlsx()
        elif fmt == 'csv':
            self.export_csv()
        elif fmt == 'txt':
            self.export_txt()
        elif fmt == 'pdf':
            self.export_pdf()

    def _download(self, filename, data, mime):
        if isinstance(data, str):
            data = data.encode('utf-8')
        ta = window.Uint8Array.new(len(data))
        ta.assign(data)
        window.saveFile(filename, ta, mime)
        self.toast(self.tr('exported', f=filename))

    def export_xlsx(self):
        sections = self.build_report_sections()
        header = {'timer_name': self.login.get('timer_name', ''),
                  'date': self.login.get('date', '')}
        wb = load_workbook(TEMPLATE_FILENAME)
        populate_template_workbook(wb, sections, header)
        bio = io.BytesIO()
        wb.save(bio)
        self._download('TimerReport.xlsx', bio.getvalue(),
                       'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

    def export_txt(self):
        sections = self.build_report_sections()
        lines = ['TIMER REPORT']
        for label, value in self.report_meta_lines():
            lines.append(f"{label}: {value}")
        lines.append('')
        for sec in SECTION_ORDER:
            headers, rows = self.section_table(sec, sections.get(sec, []))
            lines.append(self.section_title(sec))
            lines.append('\t'.join(headers))
            for r in rows:
                lines.append('\t'.join(r))
            lines.append('')
        self._download('TimerReport.txt', '\n'.join(lines), 'text/plain')

    def export_csv(self):
        sections = self.build_report_sections()
        import csv as _csv
        buf = io.StringIO()
        w = _csv.writer(buf)
        w.writerow(['TIMER REPORT'])
        for label, value in self.report_meta_lines():
            w.writerow([label, value])
        w.writerow([])
        for sec in SECTION_ORDER:
            headers, rows = self.section_table(sec, sections.get(sec, []))
            w.writerow([self.section_title(sec)])
            w.writerow(headers)
            for r in rows:
                w.writerow(r)
            w.writerow([])
        self._download('TimerReport.csv', buf.getvalue(), 'text/csv')

    def export_pdf(self):
        sections = self.build_report_sections()

        def esc(s):
            return (str(s).replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;'))

        meta_html = ''.join(
            f"<div><b>{esc(l)}:</b> {esc(v)}</div>" for l, v in self.report_meta_lines())
        body = []
        for sec in SECTION_ORDER:
            headers, rows = self.section_table(sec, sections.get(sec, []))
            body.append(f"<h3>{esc(self.section_title(sec))}</h3>")
            th = ''.join(f"<th>{esc(h)}</th>" for h in headers)
            trs = ''.join(
                '<tr>' + ''.join(f"<td>{esc(c)}</td>" for c in r) + '</tr>' for r in rows)
            body.append(f"<table><thead><tr>{th}</tr></thead><tbody>{trs}</tbody></table>")
        html = f"""<!DOCTYPE html><html><head><meta charset='utf-8'><title>Timer Report</title>
<style>
body{{font-family:Arial,'Microsoft YaHei',sans-serif;margin:24px;color:#0f172a;}}
h1{{text-align:center;margin:0 0 14px;}}
h3{{margin:18px 0 6px;font-style:italic;}}
.meta{{font-size:13px;margin-bottom:8px;}}
table{{width:100%;border-collapse:collapse;margin-bottom:6px;}}
th{{background:#1f3a8a;color:#fff;text-align:left;padding:6px;font-size:13px;}}
td{{border:1px solid #c7cdd6;padding:6px;font-size:13px;}}
</style></head><body>
<h1>TIMER REPORT</h1>
<div class='meta'>{meta_html}</div>
{''.join(body)}
</body></html>"""
        window.printReport(html)
        self.toast(self.tr('exported', f='PDF'))


_app = App()
_app.start()
